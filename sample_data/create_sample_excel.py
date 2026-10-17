"""
Generate a sample Excel file containing sales, targets, costs, and profits.
This is used to demonstrate and test all AI Excel Assistant features.
"""
import os
import random
from datetime import datetime, timedelta
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_sample_sales_data(filepath: str = "sample_data/sample_sales_data.xlsx"):
    random.seed(42)
    
    regions = ["North", "South", "East", "West"]
    categories = {
        "Hardware": ["Laptop Pro", "Cloud Tablet", "Ultra Monitor"],
        "Electronics": ["Smart Display", "Smart Hub", "Conference Cam"],
        "Accessories": ["Wireless Audio", "Ergo Keyboard", "Precision Mouse", "Fast Charger"]
    }
    
    customers = [
        "Acme Corp", "Global Tech", "Alpha Industries", "Beta Logistics", 
        "Zenith Co", "Apex Systems", "Summit Media", "Prime Retail", 
        "Vanguard Health", "Horizon Group", "Nexus Financial", "Pinnacle Energy"
    ]
    
    base_costs = {
        "Laptop Pro": (750.0, 1199.0),
        "Cloud Tablet": (320.0, 499.0),
        "Ultra Monitor": (280.0, 449.0),
        "Smart Display": (150.0, 249.0),
        "Smart Hub": (90.0, 149.0),
        "Conference Cam": (180.0, 299.0),
        "Wireless Audio": (45.0, 89.0),
        "Ergo Keyboard": (55.0, 99.0),
        "Precision Mouse": (30.0, 59.0),
        "Fast Charger": (15.0, 39.0)
    }

    start_date = datetime(2026, 1, 1)
    rows = []
    
    for i in range(1, 101):
        order_date = start_date + timedelta(days=random.randint(0, 240))
        region = random.choice(regions)
        category = random.choice(list(categories.keys()))
        product = random.choice(categories[category])
        customer = random.choice(customers)
        units = random.randint(5, 45)
        
        cost_price, list_price = base_costs[product]
        # Slight variation in pricing
        unit_cost = round(cost_price * random.uniform(0.95, 1.05), 2)
        unit_price = round(list_price * random.uniform(0.95, 1.05), 2)
        
        actual_sales = round(units * unit_price, 2)
        target_sales = round(actual_sales * random.uniform(0.85, 1.20), 2)
        total_cost = round(units * unit_cost, 2)
        profit = round(actual_sales - total_cost, 2)
        
        rows.append({
            "OrderID": f"ORD-{1000 + i}",
            "OrderDate": order_date.strftime("%Y-%m-%d"),
            "Customer": customer,
            "Region": region,
            "Category": category,
            "Product": product,
            "Units": units,
            "UnitCost": unit_cost,
            "UnitPrice": unit_price,
            "ActualSales": actual_sales,
            "TargetSales": target_sales,
            "TotalCost": total_cost,
            "Profit": profit
        })
        
    df = pd.DataFrame(rows)
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    
    # Save with openpyxl styling
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="SalesData", index=False)
        
        # Add a secondary target lookup sheet for VLOOKUP / XLOOKUP demonstrations
        targets_df = pd.DataFrame([
            {"Region": "North", "RegionalManager": "Alice Walker", "AnnualTarget": 150000, "BonusRate": 0.05},
            {"Region": "South", "RegionalManager": "Bob Martinez", "AnnualTarget": 120000, "BonusRate": 0.05},
            {"Region": "East", "RegionalManager": "Carol Smith", "AnnualTarget": 180000, "BonusRate": 0.06},
            {"Region": "West", "RegionalManager": "David Zhang", "AnnualTarget": 160000, "BonusRate": 0.055},
        ])
        targets_df.to_excel(writer, sheet_name="RegionalTargets", index=False)
        
    # Apply pretty styling to the generated workbook
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for ws in wb.worksheets:
        for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
            cell = col[0]
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = thin_border
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
                    
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    wb.save(filepath)
    print(f"Sample Excel generated successfully at: {filepath}")

if __name__ == "__main__":
    generate_sample_sales_data()
